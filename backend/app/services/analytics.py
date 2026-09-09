from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any

from bson import ObjectId

from app.core.collections import (
    ANALYTICS_DAILY_COLLECTION_NAME,
    METRICS_COLLECTION_NAME,
    SERVER_METRICS_COLLECTION_NAME,
)
from app.core.exceptions import AppError
from app.services.database_connections import monitored_connections_filter


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _day_key(value: datetime | None = None) -> str:
    return (value or _utcnow()).astimezone(timezone.utc).date().isoformat()


def _to_iso(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _storage_from_sample(sample: dict) -> dict[str, Any]:
    engine = (sample.get("meta") or {}).get("engine")
    if engine == "oracle":
        oracle = sample.get("oracle") or {}
        storage = oracle.get("storage") or {}
        tablespaces = storage.get("tablespaces") or []
        permanent = [
            item for item in tablespaces
            if str(item.get("contents") or "").upper() != "TEMPORARY"
        ]
        temporary = [
            item for item in tablespaces
            if str(item.get("contents") or "").upper() == "TEMPORARY"
        ]
        allocated = sum(int(item.get("capacity_bytes") or 0) for item in permanent)
        used = sum(int(item.get("used_bytes") or 0) for item in permanent)
        temp = sum(int(item.get("capacity_bytes") or 0) for item in temporary)
        fra = storage.get("fra") or {}
        memory = oracle.get("memory") or {}
        return {
            "database_size_bytes": allocated or None,
            "data_size_bytes": allocated or None,
            "used_size_bytes": used or None,
            "free_size_bytes": max(allocated - used, 0) if allocated else None,
            "temp_size_bytes": temp or None,
            "recovery_size_bytes": int(fra.get("used_bytes") or 0) or None,
            "recovery_capacity_bytes": int(fra.get("limit_bytes") or 0) or None,
            "sga_bytes": int(memory.get("sga_bytes") or 0) or None,
            "pga_allocated_bytes": int(memory.get("pga_allocated_bytes") or 0) or None,
            "pga_target_bytes": int(memory.get("pga_target_bytes") or 0) or None,
        }

    if engine == "sqlserver":
        sqlserver = sample.get("sqlserver") or {}
        storage = sqlserver.get("storage") or {}
        files = storage.get("files") or []
        data_size = sum(
            int(item.get("allocated_bytes") or 0)
            for item in files
            if str(item.get("file_type") or "").upper() not in {"LOG", "LOGS"}
        )
        log_size = sum(
            int(item.get("allocated_bytes") or 0)
            for item in files
            if str(item.get("file_type") or "").upper() in {"LOG", "LOGS"}
        )
        allocated = int(storage.get("allocated_bytes") or 0)
        used = storage.get("used_bytes")
        return {
            "database_size_bytes": allocated or (data_size + log_size) or None,
            "data_size_bytes": data_size or None,
            "log_size_bytes": log_size or int(sqlserver.get("log_size_bytes") or 0) or None,
            "used_size_bytes": int(used) if used is not None else None,
            "free_size_bytes": (
                max(allocated - int(used), 0)
                if allocated and used is not None
                else None
            ),
        }

    if engine == "mysql":
        mysql = sample.get("mysql") or {}
        storage = mysql.get("storage") or {}
        total = int(storage.get("total_bytes") or 0)
        return {
            "database_size_bytes": total or None,
            "data_size_bytes": int(storage.get("data_bytes") or 0) or None,
            "index_size_bytes": int(storage.get("index_bytes") or 0) or None,
        }

    return {}


def _backup_summary(latest_backup: dict | None) -> dict[str, Any]:
    if not latest_backup:
        return {
            "last_backup_at": None,
            "last_backup_status": None,
            "last_backup_kind": None,
            "last_backup_size_bytes": None,
        }

    backup_at = latest_backup.get("finished_at") or latest_backup.get("started_at")
    size = (
        latest_backup.get("backup_size_bytes")
        or latest_backup.get("output_bytes")
        or latest_backup.get("input_bytes")
    )
    return {
        "last_backup_at": backup_at,
        "last_backup_status": latest_backup.get("status"),
        "last_backup_kind": latest_backup.get("kind"),
        "last_backup_size_bytes": int(size) if size is not None else None,
    }


async def persist_database_analytics_snapshot(
    database,
    connection: dict,
    sample: dict,
    *,
    latest_backup: dict | None = None,
) -> None:
    connection_id = str(connection["_id"])
    collected_at = sample.get("collected_at") or _utcnow()
    payload: dict[str, Any] = {
        "target_type": "database",
        "target_id": connection_id,
        "day": _day_key(collected_at),
        "source": "collector",
        "connection_name": connection.get("name"),
        "engine": connection.get("engine"),
        "collected_at": collected_at,
        "status": sample.get("status"),
        "active": sample.get("active"),
        "connections": sample.get("connections"),
        "blocked": sample.get("blocked"),
        "uptime_seconds": sample.get("uptime_seconds"),
        **{key: value for key, value in _storage_from_sample(sample).items() if value is not None},
    }
    if latest_backup is not None:
        payload.update(_backup_summary(latest_backup))

    await database[ANALYTICS_DAILY_COLLECTION_NAME].update_one(
        {
            "target_type": "database",
            "target_id": connection_id,
            "day": payload["day"],
        },
        {"$set": payload, "$setOnInsert": {"created_at": _utcnow()}},
        upsert=True,
    )


async def persist_server_analytics_snapshot(database, server: dict, sample: dict) -> None:
    server_id = str(server["_id"])
    filesystems = sample.get("filesystems") or []
    disk_total = sum(int(item.get("total_bytes") or 0) for item in filesystems)
    disk_used = sum(int(item.get("used_bytes") or 0) for item in filesystems)
    memory = sample.get("memory") or {}
    collected_at = sample.get("collected_at") or _utcnow()

    payload = {
        "target_type": "server",
        "target_id": server_id,
        "day": _day_key(collected_at),
        "source": "collector",
        "server_name": server.get("name"),
        "os_family": server.get("os_family"),
        "os_version": server.get("os_version"),
        "collected_at": collected_at,
        "status": sample.get("status"),
        "uptime_seconds": sample.get("uptime_seconds"),
        "cpu_used_percent": sample.get("cpu_used_percent"),
        "memory_total_bytes": memory.get("total_bytes"),
        "memory_used_bytes": memory.get("used_bytes"),
        "memory_used_percent": memory.get("used_percent"),
        "disk_total_bytes": disk_total or None,
        "disk_used_bytes": disk_used or None,
        "disk_used_percent": (
            round(disk_used / disk_total * 100, 2) if disk_total > 0 else None
        ),
    }
    await database[ANALYTICS_DAILY_COLLECTION_NAME].update_one(
        {
            "target_type": "server",
            "target_id": server_id,
            "day": payload["day"],
        },
        {"$set": payload, "$setOnInsert": {"created_at": _utcnow()}},
        upsert=True,
    )


async def _latest_samples(database, collection_name: str, id_path: str) -> dict[str, dict]:
    pipeline = [
        {"$sort": {"collected_at": -1}},
        {"$group": {"_id": f"${id_path}", "sample": {"$first": "$$ROOT"}}},
    ]
    cursor = await database[collection_name].aggregate(pipeline)
    docs = await cursor.to_list(None)
    return {str(item["_id"]): item["sample"] for item in docs if item.get("_id")}


async def _latest_daily_by_target(database, target_type: str) -> dict[str, dict]:
    pipeline = [
        {"$match": {"target_type": target_type}},
        {"$sort": {"collected_at": -1}},
        {"$group": {"_id": "$target_id", "item": {"$first": "$$ROOT"}}},
    ]
    cursor = await database[ANALYTICS_DAILY_COLLECTION_NAME].aggregate(pipeline)
    docs = await cursor.to_list(None)
    return {str(item["_id"]): item["item"] for item in docs if item.get("_id")}


def _month_floor(months: int) -> datetime:
    now = _utcnow()
    year = now.year
    month = now.month - max(months - 1, 0)
    while month <= 0:
        month += 12
        year -= 1
    return datetime(year, month, 1, tzinfo=timezone.utc)


def _monthly_growth(docs: list[dict], names: dict[str, str]) -> list[dict]:
    latest: dict[tuple[str, str], dict] = {}
    for doc in sorted(docs, key=lambda item: item.get("collected_at") or datetime.min.replace(tzinfo=timezone.utc)):
        size = doc.get("database_size_bytes")
        if size is None:
            continue
        collected = doc.get("collected_at")
        if not isinstance(collected, datetime):
            try:
                collected = datetime.fromisoformat(str(collected).replace("Z", "+00:00"))
            except Exception:
                continue
        month = collected.strftime("%Y-%m")
        latest[(str(doc.get("target_id")), month)] = doc

    points = []
    for (target_id, month), doc in sorted(latest.items(), key=lambda item: (item[0][1], names.get(item[0][0], item[0][0]))):
        points.append(
            {
                "connection_id": target_id,
                "name": names.get(target_id, doc.get("connection_name") or target_id),
                "month": month,
                "size_bytes": int(doc.get("database_size_bytes") or 0),
                "source": doc.get("source", "collector"),
            }
        )
    return points


def _age_days(value: datetime | str | None) -> float | None:
    if value is None:
        return None
    if not isinstance(value, datetime):
        try:
            value = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except Exception:
            return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return round((_utcnow() - value).total_seconds() / 86400, 2)


async def get_database_analytics(database, *, engine: str | None = None, months: int = 12) -> dict:
    query = monitored_connections_filter()
    if engine:
        query = {"$and": [query, {"engine": engine}]}
    connections = await database.database_connections.find(query).sort("name", 1).to_list(None)
    ids = {str(item["_id"]): item for item in connections}
    latest_samples = await _latest_samples(database, METRICS_COLLECTION_NAME, "meta.connection_id")
    latest_daily = await _latest_daily_by_target(database, "database")

    items: list[dict] = []
    for connection_id, connection in ids.items():
        sample = latest_samples.get(connection_id) or {}
        daily = latest_daily.get(connection_id) or {}
        storage = _storage_from_sample(sample)

        for key in (
            "database_size_bytes", "data_size_bytes", "used_size_bytes", "free_size_bytes",
            "temp_size_bytes", "log_size_bytes", "index_size_bytes", "recovery_size_bytes",
            "recovery_capacity_bytes", "sga_bytes", "pga_allocated_bytes", "pga_target_bytes",
        ):
            if storage.get(key) is None and daily.get(key) is not None:
                storage[key] = daily.get(key)

        allocated = storage.get("database_size_bytes")
        used = storage.get("used_size_bytes")
        used_percent = (
            round(int(used) / int(allocated) * 100, 2)
            if allocated and used is not None
            else None
        )
        backup_at = daily.get("last_backup_at")
        items.append(
            {
                "connection_id": connection_id,
                "name": connection.get("name"),
                "engine": connection.get("engine"),
                "host": connection.get("host"),
                "port": connection.get("port"),
                "database": connection.get("oracle_identifier") or connection.get("database"),
                "status": sample.get("status") or "unknown",
                "checked_at": sample.get("collected_at") or sample.get("checked_at"),
                "storage_used_percent": used_percent,
                "last_backup_at": backup_at,
                "last_backup_age_days": _age_days(backup_at),
                "last_backup_status": daily.get("last_backup_status"),
                "last_backup_kind": daily.get("last_backup_kind"),
                "last_backup_size_bytes": daily.get("last_backup_size_bytes"),
                **storage,
            }
        )

    size_total = sum(int(item.get("database_size_bytes") or 0) for item in items)
    used_total = sum(int(item.get("used_size_bytes") or 0) for item in items)
    online = sum(1 for item in items if item.get("status") in {"online", "limited"})

    from_at = _month_floor(max(1, min(months, 60)))
    history_docs = await database[ANALYTICS_DAILY_COLLECTION_NAME].find(
        {
            "target_type": "database",
            "target_id": {"$in": list(ids)},
            "collected_at": {"$gte": from_at},
            "database_size_bytes": {"$ne": None},
        },
        {"_id": 0},
    ).sort("collected_at", 1).to_list(None)

    engine_counts: dict[str, int] = {}
    for item in items:
        engine_counts[item["engine"]] = engine_counts.get(item["engine"], 0) + 1

    return {
        "generated_at": _utcnow(),
        "filter_engine": engine,
        "summary": {
            "database_count": len(items),
            "online_count": online,
            "unreachable_count": max(len(items) - online, 0),
            "total_size_bytes": size_total,
            "used_size_bytes": used_total or None,
        },
        "engine_distribution": [
            {"engine": key, "count": value}
            for key, value in sorted(engine_counts.items())
        ],
        "items": items,
        "growth": _monthly_growth(history_docs, {key: value.get("name") for key, value in ids.items()}),
    }


async def get_server_analytics(database, *, os_family: str | None = None, months: int = 12) -> dict:
    query: dict[str, Any] = {"enabled": {"$ne": False}}
    if os_family:
        query["os_family"] = os_family
    servers = await database.servers.find(query).sort("name", 1).to_list(None)
    ids = {str(item["_id"]): item for item in servers}
    latest_samples = await _latest_samples(database, SERVER_METRICS_COLLECTION_NAME, "meta.server_id")
    latest_daily = await _latest_daily_by_target(database, "server")

    items: list[dict] = []
    for server_id, server in ids.items():
        sample = latest_samples.get(server_id) or {}
        daily = latest_daily.get(server_id) or {}
        memory = sample.get("memory") or {}
        filesystems = sample.get("filesystems") or []
        disk_total = sum(int(item.get("total_bytes") or 0) for item in filesystems) or daily.get("disk_total_bytes")
        disk_used = sum(int(item.get("used_bytes") or 0) for item in filesystems) or daily.get("disk_used_bytes")
        disk_percent = (
            round(int(disk_used) / int(disk_total) * 100, 2)
            if disk_total and disk_used is not None
            else daily.get("disk_used_percent")
        )
        items.append(
            {
                "server_id": server_id,
                "name": server.get("name"),
                "hostname": server.get("hostname"),
                "os_family": server.get("os_family"),
                "os_version": server.get("os_version"),
                "status": sample.get("status") or "unknown",
                "checked_at": sample.get("collected_at"),
                "cpu_used_percent": sample.get("cpu_used_percent", daily.get("cpu_used_percent")),
                "memory_total_bytes": memory.get("total_bytes", daily.get("memory_total_bytes")),
                "memory_used_bytes": memory.get("used_bytes", daily.get("memory_used_bytes")),
                "memory_used_percent": memory.get("used_percent", daily.get("memory_used_percent")),
                "disk_total_bytes": disk_total,
                "disk_used_bytes": disk_used,
                "disk_used_percent": disk_percent,
                "uptime_seconds": sample.get("uptime_seconds", daily.get("uptime_seconds")),
            }
        )

    online = sum(1 for item in items if item.get("status") in {"online", "limited"})
    os_counts: dict[str, int] = {}
    for item in items:
        key = item.get("os_family") or "other"
        os_counts[key] = os_counts.get(key, 0) + 1

    return {
        "generated_at": _utcnow(),
        "filter_os_family": os_family,
        "summary": {
            "server_count": len(items),
            "online_count": online,
            "unreachable_count": max(len(items) - online, 0),
            "total_memory_bytes": sum(int(item.get("memory_total_bytes") or 0) for item in items),
            "total_disk_bytes": sum(int(item.get("disk_total_bytes") or 0) for item in items),
            "used_disk_bytes": sum(int(item.get("disk_used_bytes") or 0) for item in items),
        },
        "os_distribution": [
            {"os_family": key, "count": value}
            for key, value in sorted(os_counts.items())
        ],
        "items": items,
    }

def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    headers = [str(value) for value in (reader.fieldnames or [])]
    return headers, [dict(row) for row in reader]


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise AppError(
            "XLSX import requires openpyxl. Install backend requirements and retry.",
            code="ANALYTICS_XLSX_DEPENDENCY_MISSING",
            status_code=500,
        ) from exc

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        first = next(rows)
    except StopIteration:
        return [], []
    headers = [str(value).strip() if value is not None else "" for value in first]
    result: list[dict[str, Any]] = []
    for row in rows:
        result.append({headers[index]: value for index, value in enumerate(row) if index < len(headers)})
    return headers, result


def parse_growth_file(filename: str, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        return _parse_csv(content)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(content)
    raise AppError(
        "Historical growth import supports CSV or XLSX files.",
        code="ANALYTICS_IMPORT_TYPE_UNSUPPORTED",
        status_code=400,
    )


def growth_import_preview(filename: str, content: bytes) -> dict:
    headers, rows = parse_growth_file(filename, content)
    if not headers:
        raise AppError("The uploaded file has no header row.", code="ANALYTICS_IMPORT_EMPTY", status_code=400)
    column_values: dict[str, list[str]] = {}
    for header in headers:
        seen: list[str] = []
        keys: set[str] = set()
        for row in rows:
            value = row.get(header)
            if value is None or str(value).strip() == "":
                continue
            text = str(value).strip()
            if text not in keys:
                keys.add(text)
                seen.append(text)
            if len(seen) >= 100:
                break
        column_values[header] = seen
    return {
        "filename": filename,
        "headers": headers,
        "row_count": len(rows),
        "preview_rows": [
            {key: _to_iso(value) for key, value in row.items()}
            for row in rows[:20]
        ],
        "column_values": column_values,
    }


def _parse_import_date(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc) if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=timezone.utc)
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%d/%m/%Y", "%Y-%m", "%b %Y", "%B %Y"):
        try:
            parsed = datetime.strptime(text, pattern)
            return parsed.replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return parsed.astimezone(timezone.utc) if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except Exception as exc:
        raise ValueError(f"Unsupported date value: {text}") from exc


def _size_bytes(value: Any, unit: str) -> int:
    number = float(str(value).replace(",", "").strip())
    multipliers = {
        "bytes": 1,
        "b": 1,
        "kb": 1024,
        "mb": 1024 ** 2,
        "gb": 1024 ** 3,
        "tb": 1024 ** 4,
    }
    key = unit.strip().lower()
    if key not in multipliers:
        raise ValueError(f"Unsupported size unit: {unit}")
    return max(int(number * multipliers[key]), 0)


async def import_database_growth(
    database,
    filename: str,
    content: bytes,
    config: dict[str, Any],
) -> dict:
    headers, rows = parse_growth_file(filename, content)
    database_column = str(config.get("database_column") or "")
    date_column = str(config.get("date_column") or "")
    size_column = str(config.get("size_column") or "")
    unit_column = str(config.get("unit_column") or "") or None
    default_unit = str(config.get("default_unit") or "GB")
    mappings = config.get("database_map") or {}

    for column in (database_column, date_column, size_column):
        if column not in headers:
            raise AppError(
                f"Import column '{column}' was not found in the file.",
                code="ANALYTICS_IMPORT_COLUMN_INVALID",
                status_code=400,
            )
    if unit_column and unit_column not in headers:
        raise AppError("Selected unit column was not found in the file.", code="ANALYTICS_IMPORT_COLUMN_INVALID", status_code=400)

    target_ids = {str(value) for value in mappings.values() if value}
    object_ids: list[ObjectId] = []
    for value in target_ids:
        try:
            object_ids.append(ObjectId(value))
        except Exception as exc:
            raise AppError("One or more database mappings are invalid.", code="ANALYTICS_IMPORT_MAPPING_INVALID", status_code=400) from exc
    connections = await database.database_connections.find({"_id": {"$in": object_ids}}).to_list(None)
    connection_map = {str(item["_id"]): item for item in connections}
    if len(connection_map) != len(target_ids):
        raise AppError("One or more mapped databases no longer exist.", code="ANALYTICS_IMPORT_MAPPING_INVALID", status_code=400)

    imported = 0
    skipped = 0
    errors: list[str] = []
    collection = database[ANALYTICS_DAILY_COLLECTION_NAME]
    for index, row in enumerate(rows, start=2):
        source_name = str(row.get(database_column) or "").strip()
        target_id = str(mappings.get(source_name) or "")
        if not source_name or not target_id:
            skipped += 1
            continue
        try:
            collected_at = _parse_import_date(row.get(date_column))
            unit = str(row.get(unit_column) or default_unit) if unit_column else default_unit
            size = _size_bytes(row.get(size_column), unit)
            connection = connection_map[target_id]
            day = _day_key(collected_at)
            existing = await collection.find_one(
                {"target_type": "database", "target_id": target_id, "day": day},
                {"source": 1},
            )
            if existing and existing.get("source") == "collector":
                skipped += 1
                continue
            await collection.update_one(
                {"target_type": "database", "target_id": target_id, "day": day},
                {
                    "$set": {
                        "target_type": "database",
                        "target_id": target_id,
                        "day": day,
                        "source": "import",
                        "connection_name": connection.get("name"),
                        "engine": connection.get("engine"),
                        "collected_at": collected_at,
                        "database_size_bytes": size,
                        "import_filename": filename,
                    },
                    "$setOnInsert": {"created_at": _utcnow()},
                },
                upsert=True,
            )
            imported += 1
        except Exception as exc:
            skipped += 1
            if len(errors) < 20:
                errors.append(f"Row {index}: {str(exc)}")

    return {"imported": imported, "skipped": skipped, "errors": errors}
